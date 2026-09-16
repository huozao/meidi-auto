from __future__ import annotations

# ================================================
# STEP CARD
# 功能: 组装 HTML/图片/Excel 并发送邮件。
# 输入: output.html, *美的*.png, 总库存*.xlsx, EMAIL_*
# 输出: 邮件发送结果（stdout）
# 上游: 050 image.py + 050 mailtxt.py
# 下游: 无（终端动作）
# ================================================

import glob
import os
import re
import smtplib
import sys
from datetime import date, datetime
from email.mime.application import MIMEApplication
from email.mime.image import MIMEImage
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText

from dotenv import load_dotenv
from openpyxl import load_workbook


def parse_recipients(raw: str) -> list[str]:
    parts = raw.replace(";", ",").split(",")
    return [item.strip() for item in parts if item.strip()]


def mask_email(value: str | None) -> str:
    if not value:
        return "(empty)"
    if "@" not in value:
        return value[:2] + "***"
    name, domain = value.split("@", 1)
    name_masked = name[0] + "*" if len(name) <= 2 else name[:2] + "***"
    return f"{name_masked}@{domain}"


def mask_secret(value: str | None) -> str:
    return "(empty)" if not value else f"***len={len(value)}"


def resolve_inventory_folder(argv: list[str]) -> str:
    default_inventory_folder = os.path.abspath(os.path.join(os.getcwd(), "data"))
    inventory_folder = argv[1] if len(argv) >= 2 else default_inventory_folder
    print(f"✅ 使用传入路径: {inventory_folder}" if len(argv) >= 2 else f"⚠️ 未传入路径，使用默认路径: {inventory_folder}")
    if not os.path.exists(inventory_folder):
        raise FileNotFoundError(f"文件夹路径不存在: {inventory_folder}")
    return inventory_folder


def pick_latest_file(folder: str, pattern: str, required: bool = True) -> str | None:
    files = glob.glob(os.path.join(folder, pattern))
    if not files:
        if required:
            raise FileNotFoundError(f"没有找到符合条件的文件: {pattern}")
        return None
    return max(files, key=os.path.getctime)


def _format_report_date(value: object) -> str:
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    text = str(value or "").strip()
    match = re.search(r"(\d{4})[年/-](\d{1,2})[月/-](\d{1,2})", text)
    if match:
        year, month, day = match.groups()
        return f"{year}-{int(month):02d}-{int(day):02d}"
    try:
        return datetime.fromisoformat(text.replace("Z", "+00:00")).strftime("%Y-%m-%d")
    except ValueError:
        return ""


def report_date_from_excel(latest_excel: str) -> str:
    """读取库存表日期，作为日报标题的业务日期。"""
    workbook = load_workbook(latest_excel, read_only=True, data_only=True)
    try:
        if "库存表" not in workbook.sheetnames:
            return ""
        return _format_report_date(workbook["库存表"]["H3"].value)
    finally:
        workbook.close()


def build_subject(latest_excel: str) -> str:
    report_date = report_date_from_excel(latest_excel)
    date_part = f"｜{report_date}" if report_date else ""
    return f"美的库存及出入库日报{date_part}｜库存、出入库及月计划"


def build_message(email_user: str, to_list: list[str], subject: str, html_content: str, latest_image: str | None, latest_excel: str) -> MIMEMultipart:
    msg = MIMEMultipart()
    msg["From"] = email_user
    msg["To"] = ", ".join(to_list)
    msg["Subject"] = subject

    body = f"""
    <html>
      <body>
        <p>您好，</p>
        <div>{html_content}</div>
        <p>祝您工作顺利！</p>
        <p>附件：<br>
          图片文件: {os.path.basename(latest_image) if latest_image else '无图片'}<br>
          Excel文件: {os.path.basename(latest_excel)}
        </p>
      </body>
    </html>
    """
    msg.attach(MIMEText(body, "html"))

    if latest_image:
        with open(latest_image, "rb") as img_file:
            img = MIMEImage(img_file.read(), name=os.path.basename(latest_image))
            msg.attach(img)

    with open(latest_excel, "rb") as excel_file:
        attachment = MIMEApplication(excel_file.read())
        attachment.add_header("Content-Disposition", "attachment", filename=os.path.basename(latest_excel))
        msg.attach(attachment)

    return msg


def send_message(email_user: str, email_password: str, msg: MIMEMultipart) -> None:
    server = smtplib.SMTP("smtp.qq.com", 587)
    server.starttls()
    server.login(email_user, email_password)
    server.send_message(msg)
    server.quit()


def main(argv: list[str] | None = None) -> int:
    argv = argv or sys.argv
    try:
        inventory_folder = resolve_inventory_folder(argv)
        latest_image = pick_latest_file(inventory_folder, "*美的*.png", required=False)
        latest_excel = pick_latest_file(inventory_folder, "*总库存*.xlsx", required=True)
        latest_html = pick_latest_file(inventory_folder, "output.html", required=True)

        with open(latest_html, "r", encoding="utf-8") as file:
            html_content = file.read()

        load_dotenv()
        email_user = os.getenv("EMAIL_ADDRESS_QQ")
        email_password = os.getenv("EMAIL_PASSWORD_QQ") or os.getenv("EMAIL_PASSWOR_QQ")
        recipient_raw = os.getenv("RECIPIENT_EMAILS", "")

        if not email_user or not email_password:
            raise ValueError("环境变量未正确配置，无法获取邮箱账户或密码")

        to_email_list = parse_recipients(recipient_raw)
        if not to_email_list:
            raise ValueError("未设置 RECIPIENT_EMAILS，无法确定收件人列表")

        print("📬 环境变量检查：")
        print("   EMAIL_ADDRESS_QQ =", mask_email(email_user))
        print("   EMAIL_PASSWORD_QQ/EMAIL_PASSWOR_QQ =", mask_secret(email_password))
        print("   RECIPIENT_EMAILS 数量 =", len(to_email_list))

        subject = build_subject(latest_excel)
        print("✉️ 邮件主题 =", subject)
        msg = build_message(email_user, to_email_list, subject, html_content, latest_image, latest_excel)
        send_message(email_user, email_password, msg)

        print("✅ 邮件发送成功！")
        return 0
    except Exception as exc:
        print(f"❌ 发送邮件失败: {exc}")
        return 1


if __name__ == "__main__":
    raise SystemExit(main())

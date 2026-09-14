# ================================================
# STEP CARD
# 功能: 生成邮件正文 HTML（异常信息与摘要）。
# 输入: 总库存*.xlsx
# 输出: output.html
# 上游: 042 Color display.py
# 下游: 051 Send an email.py
# ================================================

import os
import sys
import openpyxl
from dotenv import load_dotenv
from datetime import datetime
import re
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parent.parent
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))

from pipeline.io_utils import ensure_existing_dir, find_required_excel, resolve_data_dir
from pipeline.usage_trends import build_usage_sections


# ================================
# 📂 配置文件路径
# ================================
def get_inventory_folder():
    inventory_folder = str(resolve_data_dir(sys.argv[1] if len(sys.argv) >= 2 else None))
    print(f"✅ 使用传入路径: {inventory_folder}" if len(sys.argv) >= 2 else f"⚠️ 未传入路径，使用默认路径: {inventory_folder}")
    try:
        ensure_existing_dir(Path(inventory_folder), "库存目录")
    except FileNotFoundError as exc:
        print(f"❌ {exc}")
        sys.exit(1)

    print(f"📂 当前工作文件夹: {inventory_folder}")
    return inventory_folder


# ================================
# 1. 查找 Excel 文件
# ================================
def find_excel_file(inventory_folder):
    try:
        excel = find_required_excel(Path(inventory_folder), "总库存*.xlsx")
    except FileNotFoundError as exc:
        print(f"❌ {exc}")
        sys.exit(1)
    if os.path.basename(excel).startswith("~$"):
        print("❌ 找到的是临时锁文件，请关闭正在编辑的 Excel 后重试。")
        sys.exit(1)
    print(f"✅ 找到文件：{excel}")
    return excel


# ================================
# 2. 读取工作表
# ================================
def load_worksheet(inventory_file, sheet_name="库存表"):
    try:
        wb = openpyxl.load_workbook(inventory_file)
    except Exception as e:
        print(f"❌ 无法打开 Excel 文件：{e}")
        sys.exit(1)

    if sheet_name not in wb.sheetnames:
        print(f"❌ 工作表“{sheet_name}”不存在！")
        sys.exit(1)

    return wb[sheet_name]


# ================================
# 🎨 颜色判断函数（只识别填充色）
# ================================
def get_cell_fill_rgb(cell):
    fill = cell.fill
    if fill and fill.fill_type == "solid":
        color = fill.fgColor
        if color.type == "rgb" and color.rgb:
            return color.rgb[-6:].upper()
    return None


def is_fill_color(cell, color_code: str):
    return get_cell_fill_rgb(cell) == color_code.upper()


# ================================
# 查找红色或紫色的行
# ================================
def find_colored_rows(sheet):
    colored_rows = []
    for row in range(2, sheet.max_row + 1):
        cell = sheet.cell(row=row, column=12)
        if is_fill_color(cell, "FF0000") or is_fill_color(cell, "3F0065"):
            colored_rows.append(row)
            color_hex = get_cell_fill_rgb(cell)
            print(f"✅ 符合条件颜色 → 行: {row} RGB: #{color_hex}")
    return colored_rows


# ================================
# 📅 获取日期（H3 与 M3）
# ================================
def _fmt_dt(v):
    """把单元格时间或字符串格式化为 'YYYY-MM-DD HH:MM:SS'；无法解析就原样返回/空串。"""
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(v, str):
        s = v.strip()
        for fmt in ("%Y-%m-%d %H:%M:%S", "%Y/%m/%d %H:%M:%S", "%Y-%m-%d", "%Y/%m/%d"):
            try:
                return datetime.strptime(s, fmt).strftime("%Y-%m-%d %H:%M:%S")
            except ValueError:
                pass
        try:
            # ISO8601 兜底，如 2025-09-19T18:00:05 或带Z
            return datetime.fromisoformat(s.replace("Z", "+00:00")).strftime("%Y-%m-%d %H:%M:%S")
        except Exception:
            return s
    return ""


def get_dates(sheet):
    """
    返回 (date, date2)
    - date  来自 H3（原标题用）
    - date2 来自 M3（家里库存数据用）
    """
    date = _fmt_dt(sheet["H3"].value)
    date2 = _fmt_dt(sheet["M3"].value)
    return date, date2


def _compact_date(value):
    """日报仅展示日期，不展示具体时分秒。"""
    return value[:10] if isinstance(value, str) and len(value) >= 10 else value


# ================================
# 查找 B 列第一个空单元格
# ================================
def find_last_empty_row(sheet):
    for row in range(4, sheet.max_row + 1):
        if sheet[f"B{row}"].value is None:
            return row
    return sheet.max_row + 1


# ================================
# 计算公式并返回总和
# ================================
def calculate_sum(sheet, formula):
    if isinstance(formula, (int, float)):
        return formula
    if not isinstance(formula, str):
        return 0

    match = re.match(r"^=SUM\((.+)\)$", formula)
    if match:
        cell_range = match.group(1)
        start_cell, end_cell = cell_range.split(":")
        start_row, start_col = int(start_cell[1:]), openpyxl.utils.cell.column_index_from_string(start_cell[:1])
        end_row, end_col = int(end_cell[1:]), openpyxl.utils.cell.column_index_from_string(end_cell[:1])

        total = 0
        for row in range(start_row, end_row + 1):
            value = sheet.cell(row=row, column=start_col).value
            if isinstance(value, (int, float)):
                total += value
        return total
    return 0


# ================================
# 获取库存合计信息
# ================================
def prepare_summary_text(sheet, last_empty_row):
    stock_total       = calculate_sum(sheet, sheet.cell(row=last_empty_row, column=10).value)
    monthly_plan      = calculate_sum(sheet, sheet.cell(row=last_empty_row, column=16).value)
    plan_gap_output   = calculate_sum(sheet, sheet.cell(row=last_empty_row, column=17).value)
    monthly_sent      = calculate_sum(sheet, sheet.cell(row=last_empty_row, column=18).value)
    monthly_received  = calculate_sum(sheet, sheet.cell(row=last_empty_row, column=19).value)
    home_stock_total  = calculate_sum(sheet, sheet.cell(row=last_empty_row, column=13).value)
    monthly_remaining = monthly_plan - monthly_sent if monthly_plan and monthly_sent else 0

    print(f"📊 外仓库存: {stock_total}, 家里库存: {home_stock_total}, 月计划: {monthly_plan}, 缺口排产: {plan_gap_output}, 出库: {monthly_sent}, 入库: {monthly_received}")
    return stock_total, monthly_plan, plan_gap_output, monthly_sent, monthly_received, monthly_remaining, home_stock_total


# ================================
# 构建输出文本
# ================================
def construct_html_content(sheet, colored_rows, date, date2,
                           stock_total, monthly_plan, plan_gap_output,
                           monthly_sent, monthly_received, monthly_remaining, home_stock_total,
                           usage_sections="", auto_list_warning=""):
    html = """
    <html>
    <head>
        <meta charset="UTF-8">
        <style>
            body { margin: 0; padding: 0; background: #f3f6fa; color: #243447; font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", "Microsoft YaHei", Arial, sans-serif; font-size: 14px; line-height: 1.5; }
            body > * { box-sizing: border-box; }
            h1 { margin: 0; padding: 18px 22px 4px; color: #173b63; font-size: 22px; letter-spacing: .2px; }
            .date-grid { display: flex; gap: 18px; margin: 4px 22px 2px; color: #718096; }
            .date-card { flex: 0 0 auto; padding: 2px 0 2px 8px; border-left: 2px solid #c9dced; }
            .date-card span { color: #718096; font-size: 11px; }
            .date-card strong { margin-left: 4px; color: #526579; font-size: 12px; font-weight: 500; }
            h2 { margin: 28px 0 8px; padding: 10px 14px; border-left: 5px solid #3b82c4; background: #e8f1fb; color: #173b63; font-size: 18px; }
            h5 { margin: 6px 22px 10px; color: #526579; font-size: 13px; font-weight: 500; }
            p { margin: 8px 22px; color: #5a6b7d; }
            table { border-collapse: separate; border-spacing: 0; width: 100%; min-width: 560px; margin: 0; background: #fff; }
            th, td { border-right: 1px solid #d9e2ec; border-bottom: 1px solid #e4eaf1; padding: 8px 10px; white-space: nowrap; }
            th { background: #eaf2fb; color: #173b63; text-align: left; font-weight: 650; }
            tbody tr:nth-child(even) { background: #f8fbff; }
            tbody tr.total-row { background: #dcecfb; color: #173b63; font-weight: 700; }
            tbody tr.total-row td { border-top: 2px solid #9bb9d6; }
            tbody tr:hover { background: #fff5d6; }
            td.right, td.num { text-align: right; font-variant-numeric: tabular-nums; }
            td.left { text-align: left; }
            td.code, td.unit, td.rank { text-align: center; }
            td.today-change { color: #16804b; font-weight: 700; text-align: right; }
            .table-wrap { overflow-x: auto; margin: 10px 22px 24px; border: 1px solid #d7e1ec; border-radius: 10px; box-shadow: 0 2px 8px rgba(31, 58, 95, .06); }
            .table-wrap table tr:first-child th:first-child { border-top-left-radius: 9px; }
            .table-wrap table tr:first-child th:last-child { border-top-right-radius: 9px; }
            .summary-table { min-width: 420px; }
            .trend { min-width: 760px; table-layout: fixed; }
            .company { min-width: 760px; table-layout: fixed; }
            .trend .code-col, .company .code-col { width: 66px; }
            .trend .month-col, .company .month-col { width: 86px; }
            .trend .delta-col, .company .delta-col { width: 88px; }
            .trend .unit-col, .company .unit-col { width: 58px; }
            .company .rank-col { width: 48px; }
            .company .company-col { width: 96px; }
            .company .total-col { width: 96px; }
            .company-name { white-space: normal; word-break: break-all; }
            .total-label { text-align: left; }
            .warn { color: #a61c00; background: #fff1f0; border: 1px solid #f3b5ae; border-radius: 6px; padding: 8px 12px; }
            @media (max-width: 700px) { h1 { font-size: 19px; } .table-wrap { margin-left: 10px; margin-right: 10px; } p, h5 { margin-left: 10px; margin-right: 10px; } }
        </style>
    </head>
    <body>
    """

    # 两个数据时间卡片：H3 对应重庆俊都仓储，M3 对应家里库存
    html += f"""
    <h1>美的仓储日报</h1>
    <div class="date-grid">
        <div class="date-card"><span>俊都仓储</span><strong>{_compact_date(date)}</strong></div>
        <div class="date-card"><span>家里库存</span><strong>{_compact_date(date2)}</strong></div>
    </div>
    <h5>库存预警物料有 <strong>{len(colored_rows)}</strong> 款</h5>
    """
    if auto_list_warning:
        html += f'<p class="warn">{auto_list_warning}</p>'

    html += """
    <div class="table-wrap"><table class="summary-table">
        <tr>
            <th>编号</th>
            <th>库存</th>
            <th>外应存（3月周均）</th>
            <th>家里库存</th>
        </tr>
    """
    for row in colored_rows:
        code = sheet.cell(row=row, column=3).value
        stock = sheet.cell(row=row, column=10).value
        expected = sheet.cell(row=row, column=11).value
        home_stock = sheet.cell(row=row, column=13).value

        stock_fmt = f"{stock:,.1f}" if isinstance(stock, (int, float)) else stock
        expected_fmt = f"{expected:,.1f}" if isinstance(expected, (int, float)) else expected
        home_stock_fmt = f"{home_stock:,.1f}" if isinstance(home_stock, (int, float)) else home_stock

        html += f"""
        <tr>
            <td>{code}</td>
            <td class="right">{stock_fmt}</td>
            <td class="right">{expected_fmt}</td>
            <td class="right">{home_stock_fmt}</td>
        </tr>
        """

    html += "</table></div>"
    html += """
    <h5>汇总信息</h5>
    <div class="table-wrap"><table class="summary-table">
        <tr><th>库存项目</th><th>数值</th><th>计划项目</th><th>数值</th></tr>
    """

    def cell(label, value):
        return f'<td class="left">{label}</td><td class="right">{value:,.1f}</td>'

    html += "<tr>" + cell("出库", monthly_sent) + cell("月计划", monthly_plan) + "</tr>"
    html += "<tr>" + cell("入库", monthly_received) + cell("月计划预估还有要发货", monthly_remaining) + "</tr>"
    html += "<tr>" + cell("外仓库存总量", stock_total) + cell("月计划缺口排产", plan_gap_output) + "</tr>"
    html += "<tr>" + cell("家里库存总量", home_stock_total) + '<td class="left"></td><td class="right"></td></tr>'

    html += "</table></div>"
    if usage_sections:
        html += usage_sections
    html += "\n</body></html>"
    return html


# ================================
# 保存为 HTML 文件
# ================================
def save_output_to_file(html_content, output_dir):
    output_filename = os.path.join(output_dir, "output.html")
    with open(output_filename, "w", encoding="utf-8") as f:
        f.write(html_content)
    print(f"📁 已成功保存为 HTML 文件：{output_filename}")


# ================================
# 主函数
# ================================
def main(argv: list[str] | None = None) -> int:
    if argv is not None:
        sys.argv = argv
    load_dotenv(REPO_ROOT / ".env")
    inventory_folder = get_inventory_folder()
    inventory_file = find_excel_file(inventory_folder)
    sheet = load_worksheet(inventory_file)

    colored_rows = find_colored_rows(sheet)
    date, date2 = get_dates(sheet)  # 👈 同时拿 H3 / M3
    last_empty_row = find_last_empty_row(sheet)
    print(f"⚡ 发现 B 列第一个空单元格所在行: {last_empty_row}")

    stock_total, monthly_plan, plan_gap_output, monthly_sent, monthly_received, monthly_remaining, home_stock_total = prepare_summary_text(sheet, last_empty_row)
    monthly_root_value = os.getenv("MONTHLY_ARCHIVE_DIR", "").strip()
    usage_sections = build_usage_sections(
        Path(inventory_folder),
        Path(monthly_root_value).expanduser() if monthly_root_value else None,
    )
    auto_list_warning = ""
    if (Path(inventory_folder) / ".auto-list-unavailable").exists():
        auto_list_warning = "未自动获取到最新自动清单，无法计算外应存、家应存和月计划；本邮件仅展示原始库存及趋势数据。"

    html_content = construct_html_content(
        sheet, colored_rows, date, date2,
        stock_total, monthly_plan, plan_gap_output,
        monthly_sent, monthly_received, monthly_remaining, home_stock_total, usage_sections, auto_list_warning
    )

    print("\n📋 HTML 已生成，预览内容省略…")
    save_output_to_file(html_content, inventory_folder)
    print("✅ 红色或紫色单元格数量：", len(colored_rows))
    print("📌 行号列表：", colored_rows)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

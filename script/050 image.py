from __future__ import annotations

# ================================================
# STEP CARD
# 功能: 按 Excel 原生分页和样式导出库存表邮件图片附件。
# 输入: 总库存*.xlsx
# 输出: *美的*.png
# 上游: 042 Color display.py
# 下游: 051 Send an email.py
# 运行依赖: LibreOffice Calc、中文字体、poppler-utils
# ================================================

import glob
import os
import re
import shutil
import subprocess
import sys
import tempfile
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter
from PIL import Image, ImageChops


SHEET_NAME = "库存表"
DEFAULT_COL_RANGE = "A:T"
RENDER_DPI = 150
MAX_IMAGE_WIDTH = 1800


def resolve_inventory_folder(argv: list[str] | None = None) -> str:
    default_inventory_folder = os.path.abspath(os.path.join(os.getcwd(), "data"))
    argv = argv or sys.argv

    positional = [item for item in argv[1:] if not item.startswith("-")]
    if positional:
        inventory_folder = positional[-1]
        print(f"✅ 使用传入路径: {inventory_folder}")
    else:
        inventory_folder = default_inventory_folder
        print(f"⚠️ 未传入路径，使用默认路径: {inventory_folder}")

    if not os.path.exists(inventory_folder):
        print(f"❌ 文件夹路径不存在: {inventory_folder}")
        raise SystemExit(1)
    return inventory_folder


def pick_inventory_file(inventory_folder: str) -> str:
    files = glob.glob(os.path.join(inventory_folder, "总库存*.xlsx"))
    files = [f for f in files if not os.path.basename(f).startswith("~$")]
    if not files:
        print("❌ 没有找到符合条件的总库存文件！")
        raise SystemExit(1)
    latest_file = max(files, key=os.path.getctime)
    print(f"✅ 找到最新文件：{latest_file}")
    return latest_file


def _is_non_empty(value: object) -> bool:
    return value is not None and (not isinstance(value, str) or bool(value.strip()))


def _parse_col_range(value: str) -> tuple[int, int]:
    """解析 A:T / A-T / A~T 形式。"""
    text = value.strip().upper().replace(" ", "")
    match = re.fullmatch(r"([A-Z]+)[:\-~]([A-Z]+)", text)
    if not match:
        raise ValueError(f"列范围格式错误: {value}（示例: A:T）")
    first = column_index_from_string(match.group(1))
    last = column_index_from_string(match.group(2))
    return min(first, last), max(first, last)


def detect_used_bounds(ws, col_start: int, col_end: int) -> tuple[int, int]:
    """在指定列范围内检测首尾有效行，并保留少量上下边距。"""
    populated_rows = [
        row
        for row in range(1, ws.max_row + 1)
        if any(_is_non_empty(ws.cell(row=row, column=col).value) for col in range(col_start, col_end + 1))
    ]
    if not populated_rows:
        return 1, min(ws.max_row, 40)
    return max(1, min(populated_rows) - 1), min(ws.max_row, max(populated_rows) + 2)


def _prepare_render_workbook(
    source_path: Path,
    render_path: Path,
    *,
    col_start: int,
    col_end: int,
) -> tuple[int, int]:
    """复制工作簿并只保留库存表的邮件预览区域，不改动原始 Excel。"""
    workbook = load_workbook(source_path, data_only=False)
    try:
        if SHEET_NAME not in workbook.sheetnames:
            raise ValueError(f"工作簿中不存在工作表: {SHEET_NAME}")
        worksheet = workbook[SHEET_NAME]
        row_start, row_end = detect_used_bounds(worksheet, col_start, col_end)
        area = f"{get_column_letter(col_start)}{row_start}:{get_column_letter(col_end)}{row_end}"

        worksheet.print_area = area
        worksheet.sheet_properties.pageSetUpPr.fitToPage = True
        worksheet.page_setup.orientation = worksheet.ORIENTATION_LANDSCAPE
        worksheet.page_setup.paperSize = worksheet.PAPERSIZE_A3
        worksheet.page_setup.fitToWidth = 1
        worksheet.page_setup.fitToHeight = 1
        worksheet.page_margins.left = 0.2
        worksheet.page_margins.right = 0.2
        worksheet.page_margins.top = 0.25
        worksheet.page_margins.bottom = 0.25
        worksheet.sheet_view.showGridLines = True

        for sheet in workbook.worksheets:
            if sheet.title != SHEET_NAME:
                sheet.sheet_state = "hidden"
        workbook.active = workbook.index(worksheet)
        if workbook.calculation is not None:
            workbook.calculation.fullCalcOnLoad = True
            workbook.calculation.forceFullCalc = True
        workbook.save(render_path)
        return row_start, row_end
    finally:
        workbook.close()


def _command_path(*names: str) -> str | None:
    for name in names:
        resolved = shutil.which(name)
        if resolved:
            return resolved
    return None


def render_excel_preview(source_path: str | Path, output_path: str | Path, col_range: str = DEFAULT_COL_RANGE) -> tuple[int, int]:
    """使用 LibreOffice + poppler 将库存表原生渲染为 PNG。"""
    source = Path(source_path)
    output = Path(output_path)
    soffice = _command_path("soffice", "libreoffice")
    pdftoppm = _command_path("pdftoppm")
    if not soffice or not pdftoppm:
        missing = []
        if not soffice:
            missing.append("LibreOffice (soffice)")
        if not pdftoppm:
            missing.append("poppler-utils (pdftoppm)")
        raise RuntimeError("邮件图片原生渲染缺少: " + ", ".join(missing))

    col_start, col_end = _parse_col_range(col_range)
    output.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="meidi-excel-preview-") as temp_dir:
        temp = Path(temp_dir)
        render_workbook = temp / source.name
        row_start, row_end = _prepare_render_workbook(
            source,
            render_workbook,
            col_start=col_start,
            col_end=col_end,
        )

        profile = temp / "lo-profile"
        profile_uri = profile.resolve().as_uri()
        env = os.environ.copy()
        env["HOME"] = str(temp / "home")
        env["SAL_USE_VCLPLUGIN"] = "headless"
        pdf_dir = temp / "pdf"
        pdf_dir.mkdir()
        subprocess.run(
            [
                soffice,
                f"-env:UserInstallation={profile_uri}",
                "--headless",
                "--convert-to",
                "pdf:calc_pdf_Export",
                "--outdir",
                str(pdf_dir),
                str(render_workbook),
            ],
            check=True,
            capture_output=True,
            text=True,
            timeout=120,
            env=env,
        )
        pdf_path = pdf_dir / f"{render_workbook.stem}.pdf"
        if not pdf_path.exists():
            raise RuntimeError(f"LibreOffice 未生成 PDF: {pdf_path}")

        png_prefix = temp / "inventory-preview"
        subprocess.run(
            [
                pdftoppm,
                "-png",
                "-r",
                str(RENDER_DPI),
                "-singlefile",
                str(pdf_path),
                str(png_prefix),
            ],
            check=True,
            capture_output=True,
            text=True,
            timeout=120,
        )
        rendered_png = png_prefix.with_suffix(".png")
        if not rendered_png.exists():
            raise RuntimeError(f"poppler 未生成 PNG: {rendered_png}")
        with Image.open(rendered_png) as rendered:
            image = rendered.convert("RGB")
            background = Image.new("RGB", image.size, "white")
            bounds = ImageChops.difference(image, background).getbbox()
            if bounds:
                padding = 24
                left = max(0, bounds[0] - padding)
                top = max(0, bounds[1] - padding)
                right = min(image.width, bounds[2] + padding)
                bottom = min(image.height, bounds[3] + padding)
                image = image.crop((left, top, right, bottom))
            if image.width > MAX_IMAGE_WIDTH:
                height = round(image.height * MAX_IMAGE_WIDTH / image.width)
                image = image.resize((MAX_IMAGE_WIDTH, height), Image.Resampling.LANCZOS)
            image.save(output, format="PNG", optimize=True)

    print(f"🖼️ Excel 原生预览区域: {get_column_letter(col_start)}{row_start}:{get_column_letter(col_end)}{row_end}")
    return row_start, row_end


def main(argv: list[str] | None = None) -> int:
    folder = resolve_inventory_folder(argv or sys.argv)
    latest_file = pick_inventory_file(folder)
    col_range = os.getenv("MAIL_IMAGE_COL_RANGE", DEFAULT_COL_RANGE)
    current_time = datetime.now().strftime("%Y%m%d_%H%M%S")
    image_filepath = os.path.join(folder, f"美的仓储自动化_{current_time}.png")
    render_excel_preview(latest_file, image_filepath, col_range=col_range)
    print(f"✅ 图片已保存：{image_filepath}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

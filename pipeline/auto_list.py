"""从最近三个完整月份的出库数据生成每日库存流水线使用的自动清单。"""

from __future__ import annotations

import os
import tempfile
from dataclasses import dataclass
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill

from .monthly_summary import AggregationResult, month_token, normalize_code


AUTO_LIST_HEADERS = ("编号", "外应存(3月2周均)", "家应存(3月2周均)", "月计划(3月月均)", "备注", "比例", "近3月出库总量", "基准月份")


@dataclass(frozen=True)
class AutoListResult:
    path: Path
    anchor_month: str
    source_months: tuple[str, ...]
    material_count: int


def _previous_month(value: str) -> str:
    year, month = map(int, value.split("-"))
    return f"{year - 1:04d}-12" if month == 1 else f"{year:04d}-{month - 1:02d}"


def rolling_months(anchor_month: str, count: int = 3) -> tuple[str, ...]:
    months = [anchor_month]
    for _ in range(count - 1):
        months.append(_previous_month(months[-1]))
    return tuple(reversed(months))


def _short_code(value: object) -> str:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ""
    code = normalize_code(value)
    return code.zfill(5) if code.isdigit() and len(code) <= 5 else code


def load_remarks(path: Path) -> dict[str, object]:
    """只继承人工维护的备注列，其他数量字段每月重新计算。"""
    if not path.exists():
        return {}
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.active
        headers = {str(cell.value).strip(): index for index, cell in enumerate(next(ws.iter_rows(min_row=1, max_row=1, values_only=False)), start=1) if cell.value}
        code_col, note_col = headers.get("编号"), headers.get("备注")
        if not code_col or not note_col:
            return {}
        remarks: dict[str, object] = {}
        for row in ws.iter_rows(min_row=2, values_only=True):
            code = _short_code(row[code_col - 1] if code_col <= len(row) else None)
            if code:
                remarks[code] = row[note_col - 1] if note_col <= len(row) else None
        return remarks
    finally:
        wb.close()


def build_auto_list_frame(result: AggregationResult, anchor_month: str) -> tuple[pd.DataFrame, tuple[str, ...]]:
    source_months = rolling_months(anchor_month)
    records = pd.DataFrame(result.records)
    available = set(records.get("month", pd.Series(dtype=str)).astype(str))
    missing = [month for month in source_months if month not in available]
    if missing:
        raise ValueError(f"无法生成自动清单，缺少近三个月有效月末数据: {', '.join(missing)}")

    window = records[records["month"].isin(source_months)].copy()
    window["编号"] = window["short_code"].map(_short_code)
    window = window[window["编号"] != ""]
    if window.empty:
        raise ValueError("无法生成自动清单：近三个月没有可用物料编码")

    latest = window[window["month"] == anchor_month].copy()
    latest = latest.sort_values(["编号", "code"]).drop_duplicates("编号", keep="last")
    totals = window.groupby("编号", as_index=False)["outbound"].sum().rename(columns={"outbound": "近3月出库总量"})
    frame = latest[["编号"]].merge(totals, on="编号", how="left")
    frame["近3月出库总量"] = frame["近3月出库总量"].fillna(0.0)
    # 近三个月总量 / 6 = 两周用量；外仓和家里各承担两周，合计覆盖一个月。
    frame["外应存"] = (frame["近3月出库总量"] / 6).round(2)
    frame["家应存"] = (frame["近3月出库总量"] / 6).round(2)
    frame["月计划"] = (frame["近3月出库总量"] / 3).round(2)
    frame["备注"] = None
    frame["比例"] = None
    frame["基准月份"] = f"{month_token(source_months[0])}~{month_token(source_months[-1])}"
    frame = frame.rename(columns={
        "外应存": "外应存(3月2周均)",
        "家应存": "家应存(3月2周均)",
        "月计划": "月计划(3月月均)",
    })
    return frame.loc[:, AUTO_LIST_HEADERS].sort_values("编号").reset_index(drop=True), source_months


def write_auto_list(frame: pd.DataFrame, output_path: Path, anchor_month: str) -> Path:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = month_token(anchor_month)
    ws.append(list(AUTO_LIST_HEADERS))
    for row in frame.itertuples(index=False, name=None):
        ws.append(list(row))
    header_fill = PatternFill(start_color="D9EAD3", end_color="D9EAD3", fill_type="solid")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    for column in range(2, 5):
        for cell in ws.iter_cols(min_col=column, max_col=column, min_row=2):
            for value in cell:
                value.number_format = "0.00"
    for letter, width in {"A": 12, "B": 12, "C": 12, "D": 12, "E": 28, "F": 10, "G": 16, "H": 14}.items():
        ws.column_dimensions[letter].width = width
    ws.freeze_panes = "A2"
    with tempfile.NamedTemporaryFile(dir=output_path.parent, prefix=".auto-list-", suffix=".xlsx", delete=False) as temp:
        temp_path = Path(temp.name)
    try:
        wb.save(temp_path)
        os.replace(temp_path, output_path)
    finally:
        wb.close()
        if temp_path.exists():
            temp_path.unlink()
    return output_path


def generate_auto_list(
    result: AggregationResult,
    anchor_month: str,
    output_path: Path,
    remarks_source: Path | None = None,
) -> AutoListResult:
    frame, source_months = build_auto_list_frame(result, anchor_month)
    remarks = load_remarks(remarks_source or output_path)
    frame["备注"] = frame["编号"].map(remarks)
    write_auto_list(frame, output_path, anchor_month)
    return AutoListResult(output_path, anchor_month, source_months, len(frame))

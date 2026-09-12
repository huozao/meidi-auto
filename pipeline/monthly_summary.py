"""月末邮件选择、月度工作簿扫描与物料汇总。

该模块不参与每日生产流水线，供 ``tools/monthly_summary.py`` 调用。
所有时间统一使用 Asia/Shanghai；输入文件只读，自动归档和报告使用独立目录。
"""

from __future__ import annotations

import calendar
import fnmatch
import hashlib
import imaplib
import json
import os
import re
import tempfile
from collections import OrderedDict
from dataclasses import dataclass
from datetime import date, datetime, time, timedelta
from email import message_from_bytes
from email.header import decode_header, make_header
from email.utils import parsedate_to_datetime
from pathlib import Path
from typing import Iterable, Sequence
from zoneinfo import ZoneInfo

import openpyxl
import pandas as pd
from openpyxl.chart import BarChart, LineChart, Reference


TZ_SHANGHAI = ZoneInfo("Asia/Shanghai")
MONTH_FILE_RE = re.compile(r"^(?P<token>\d{4})月底\.xlsx$")
MONTH_RE = re.compile(r"^(?P<year>\d{4})-(?P<month>0[1-9]|1[0-2])$")

REQUIRED_WORKBOOK_COLUMNS = ("外仓入库总量", "外仓出库总量", "库存")
OPTIONAL_WORKBOOK_COLUMNS = ("合格仓库存", "外应存", "月计划")

COMPANY_ALIASES = (
    ("重庆渝丰和科技有限公司", ("重庆渝丰和科技有限公司", "渝丰和")),
    ("重庆恒讯联盛实业有限公司", ("重庆恒讯联盛实业有限公司", "恒讯")),
    ("重庆郅塑科技有限公司", ("重庆郅塑科技有限公司", "郅塑", "致塑")),
    ("重庆松亚美迪电子有限公司", ("重庆松亚美迪电子有限公司", "松亚", "松垭")),
    ("重庆瀚海塑胶制品有限公司", ("重庆瀚海塑胶制品有限公司", "瀚海")),
    ("重庆联翼森科技有限公司", ("重庆联翼森科技有限公司", "联翼森")),
    ("重庆迅飞汽车零部件有限公司", ("重庆迅飞汽车零部件有限公司", "讯飞")),
    ("重庆欧盼科技发展有限公司", ("重庆欧盼科技发展有限公司", "欧盼")),
    ("重庆西雄机车零部件有限公司", ("重庆西雄机车零部件有限公司", "西雄")),
    ("四川润光科技发展有限公司", ("四川润光科技发展有限公司", "润光")),
    ("重庆汉美实业有限公司", ("重庆汉美实业有限公司", "汉美")),
    ("广东佳适新材料科技有限公司", ("广东佳适新材料科技有限公司", "佳适")),
    ("重庆昊泰塑胶制品有限公司", ("重庆昊泰塑胶制品有限公司", "昊泰")),
    ("重庆赋域电子科技有限公司", ("重庆赋域电子科技有限公司", "赋域")),
    ("重庆工厂", ("重庆工厂",)),
    ("美的", ("美的",)),
)


@dataclass(frozen=True)
class MonthlySource:
    month: str
    path: Path
    source: str


@dataclass(frozen=True)
class MailAttachmentCandidate:
    uid: str
    received_at: datetime
    subject: str
    filename: str
    payload: bytes
    html: str = ""


@dataclass
class WorkbookReadResult:
    month: str
    path: Path
    records: list[dict]
    quality: dict
    movement_records: list[dict]


@dataclass
class AggregationResult:
    records: list[dict]
    quality: list[dict]
    sources: list[MonthlySource]
    movement_records: list[dict]


def normalize_text(value: object) -> str:
    if value is None:
        return ""
    return str(value).replace("\u3000", " ").strip()


def normalize_code(value: object) -> str:
    if value is None or normalize_text(value) == "":
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, int):
        return str(value)
    return normalize_text(value)


def extract_company(note: object) -> str:
    """从备注提取统一公司名；保留无法可靠归属的备注原文。"""
    text = normalize_text(note)
    if not text:
        return "未备注"
    for canonical, aliases in COMPANY_ALIASES:
        if any(alias in text for alias in aliases):
            return canonical
    if text.endswith("-材料"):
        return text[:-3].strip() or "未标注"
    if "厂家" in text:
        return "厂家"
    if "车间" in text or "无需求" in text:
        return "内部调整"
    if text.startswith("调拨"):
        return text[2:].strip() or "内部调拨"
    return text


def classify_movement(category: object, note: object, inbound: float, outbound: float) -> str:
    """将 ERP 库存变动分类为可读的业务类型。"""
    category_text = normalize_text(category)
    note_text = normalize_text(note)
    if "借" in note_text and any(token in note_text for token in ("退", "换")):
        return "借用/退换混合"
    if "借" in note_text and outbound:
        return "借用出库"
    if category_text == "出库":
        return "领用出库"
    if category_text == "入库":
        if "还" in note_text:
            return "借用归还"
        if any(token in note_text for token in ("退货", "换货")):
            return "退货/换货入库"
        return "普通入库"
    if "调出" in category_text:
        return "不合格品调出"
    if "调入" in category_text:
        return "不合格品调入"
    if "退货" in category_text or any(token in note_text for token in ("退货", "换货", "退回厂家")):
        return "退货/换货"
    return category_text or "其他"


def _read_movement_rows(wb: openpyxl.Workbook, month: str) -> list[dict]:
    if "出入库明细表" not in wb.sheetnames:
        return []
    ws = wb["出入库明细表"]
    try:
        header_row, headers = _find_header(ws, {"库存变动类别"})
    except ValueError:
        return []
    code_column = headers.get("美的编码", headers.get("代编码"))
    name_column = headers.get("物料品名")
    date_column = headers.get("出入库日期", headers.get("录入日期"))
    if code_column is None or name_column is None or date_column is None:
        return []
    year = parse_month(month)[0]
    movement_records: list[dict] = []
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        code = normalize_code(row[code_column] if code_column < len(row) else None)
        name = normalize_text(row[name_column] if name_column < len(row) else None)
        parsed_date = parse_datetime(row[date_column] if date_column < len(row) else None)
        if not code or parsed_date is None:
            continue
        category = normalize_text(row[headers["库存变动类别"]] if headers["库存变动类别"] < len(row) else None)
        inbound = to_number(row[headers["本期收入"]]) if "本期收入" in headers and headers["本期收入"] < len(row) else 0.0
        outbound = to_number(row[headers["本期发出"]]) if "本期发出" in headers and headers["本期发出"] < len(row) else 0.0
        note = normalize_text(row[headers["备注"]]) if "备注" in headers and headers["备注"] < len(row) else ""
        movement_records.append({
            "month": month,
            "year": year,
            "date": parsed_date.date().isoformat(),
            "code": code,
            "name": name,
            "unit": normalize_text(row[headers["单位"]]) if "单位" in headers and headers["单位"] < len(row) else "",
            "category": category,
            "business_type": classify_movement(category, note, inbound, outbound),
            "company": extract_company(note),
            "note": note,
            "inbound": inbound,
            "outbound": outbound,
            "quantity": outbound if outbound else inbound,
        })
    return movement_records


def to_number(value: object) -> float:
    if value is None or value == "":
        return 0.0
    if isinstance(value, bool):
        return float(value)
    if isinstance(value, (int, float)):
        return float(value)
    text = normalize_text(value).replace(",", "")
    if text in {"-", "/", "--", "无", "None", "null"}:
        return 0.0
    try:
        return float(text)
    except ValueError:
        return 0.0


def parse_month(value: str) -> tuple[int, int]:
    match = MONTH_RE.fullmatch(value.strip())
    if not match:
        raise ValueError(f"月份必须是 YYYY-MM: {value}")
    return int(match.group("year")), int(match.group("month"))


def month_token(month: str) -> str:
    year, month_number = parse_month(month)
    return f"{year % 100:02d}{month_number:02d}"


def month_cutoff(month: str) -> datetime:
    year, month_number = parse_month(month)
    last_day = calendar.monthrange(year, month_number)[1]
    return datetime.combine(date(year, month_number, last_day), time.max, tzinfo=TZ_SHANGHAI)


def month_start(month: str) -> date:
    year, month_number = parse_month(month)
    return date(year, month_number, 1)


def next_month_start(month: str) -> date:
    year, month_number = parse_month(month)
    if month_number == 12:
        return date(year + 1, 1, 1)
    return date(year, month_number + 1, 1)


def parse_datetime(value: object) -> datetime | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        parsed = value
    elif isinstance(value, date):
        parsed = datetime.combine(value, time.min)
    else:
        text = normalize_text(value)
        parsed = None
        for fmt in (
            "%Y/%m/%d %H:%M:%S",
            "%Y-%m-%d %H:%M:%S",
            "%Y年%m月%d日 %H:%M:%S",
            "%Y/%m/%d",
            "%Y-%m-%d",
            "%Y年%m月%d日",
        ):
            try:
                parsed = datetime.strptime(text, fmt)
                break
            except ValueError:
                continue
        if parsed is None:
            try:
                parsed = datetime.fromisoformat(text.replace("Z", "+00:00"))
            except ValueError:
                return None
    if parsed.tzinfo is None:
        return parsed.replace(tzinfo=TZ_SHANGHAI)
    return parsed.astimezone(TZ_SHANGHAI)


def parse_mail_datetime(value: str | None) -> datetime | None:
    if not value:
        return None
    try:
        parsed = parsedate_to_datetime(value)
    except (TypeError, ValueError, IndexError):
        return None
    if parsed.tzinfo is None:
        parsed = parsed.replace(tzinfo=TZ_SHANGHAI)
    return parsed.astimezone(TZ_SHANGHAI)


def parse_imap_internaldate(value: str | bytes | None) -> datetime | None:
    if isinstance(value, bytes):
        value = value.decode("ascii", errors="ignore")
    if not value:
        return None
    match = re.search(r'INTERNALDATE\s+"([^"]+)"', value, flags=re.IGNORECASE)
    if not match:
        return None
    try:
        parsed = datetime.strptime(match.group(1), "%d-%b-%Y %H:%M:%S %z")
    except ValueError:
        return None
    return parsed.astimezone(TZ_SHANGHAI)


def decode_mime_header(value: str | None) -> str:
    if not value:
        return ""
    try:
        return str(make_header(decode_header(value)))
    except (ValueError, UnicodeError):
        return value


def _files_in(directory: Path) -> dict[str, Path]:
    if not directory.exists():
        return {}
    found: dict[str, Path] = {}
    for path in sorted(directory.iterdir()):
        if not path.is_file():
            continue
        match = MONTH_FILE_RE.fullmatch(path.name)
        if not match:
            continue
        token = match.group("token")
        if token in found:
            raise ValueError(f"同一目录存在重复月份文件: {found[token]} / {path}")
        found[token] = path
    return found


def scan_month_files(root: Path, auto_archive_dir: Path | None = None) -> list[MonthlySource]:
    """扫描人工月末文件和自动归档文件，自动归档优先。"""
    manual = _files_in(root)
    automatic = _files_in(auto_archive_dir) if auto_archive_dir else {}
    merged = dict(manual)
    merged.update(automatic)
    sources: list[MonthlySource] = []
    for token, path in sorted(merged.items()):
        year = 2000 + int(token[:2])
        month_number = int(token[2:])
        if not 1 <= month_number <= 12:
            raise ValueError(f"非法月份文件名: {path.name}")
        sources.append(
            MonthlySource(
                month=f"{year:04d}-{month_number:02d}",
                path=path,
                source="automatic" if token in automatic else "manual",
            )
        )
    return sources


def _find_header(ws, required: set[str]) -> tuple[int, dict[str, int]]:
    for row_number, row in enumerate(ws.iter_rows(min_row=1, max_row=min(ws.max_row, 10), values_only=True), start=1):
        headers = {normalize_text(value): index for index, value in enumerate(row) if normalize_text(value)}
        if required.issubset(headers) or ("编号" in headers and required - {"物料编码"}.issubset(headers)):
            return row_number, headers
    raise ValueError(f"工作表 {ws.title} 未找到所需表头")


def _read_detail_quality(wb: openpyxl.Workbook, month: str) -> dict:
    quality = {
        "detail_sheet": "出入库明细表" in wb.sheetnames,
        "detail_last_date": "",
        "month_end_rows": 0,
        "future_rows": 0,
        "detail_warning": "",
    }
    if "出入库明细表" not in wb.sheetnames:
        quality["detail_warning"] = "缺少出入库明细表"
        return quality

    ws = wb["出入库明细表"]
    try:
        header_row, headers = _find_header(ws, {"库存变动类别"})
    except ValueError:
        quality["detail_warning"] = "出入库明细表未找到表头"
        return quality
    date_column = headers.get("出入库日期", headers.get("录入日期"))
    if date_column is None:
        quality["detail_warning"] = "出入库明细表缺少日期列"
        return quality

    cutoff_date = month_cutoff(month).date()
    parsed_dates: list[date] = []
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        parsed = parse_datetime(row[date_column] if date_column < len(row) else None)
        if parsed is None:
            continue
        parsed_date = parsed.date()
        parsed_dates.append(parsed_date)
        if parsed_date == cutoff_date:
            quality["month_end_rows"] += 1
        if parsed_date > cutoff_date:
            quality["future_rows"] += 1
    if parsed_dates:
        quality["detail_last_date"] = max(parsed_dates).isoformat()
    if quality["month_end_rows"] == 0:
        quality["detail_warning"] = "未发现月末当天出入库记录"
    if quality["future_rows"]:
        quality["detail_warning"] = "存在晚于月末的出入库记录"
    return quality


def _read_inventory_header_quality(ws, header_row: int, month: str) -> dict:
    """以库存表表头上方的日期列确认工作簿是否为目标月末快照。

    业务表可能没有月末当天的出入库流水，但库存表仍会生成当天的库存列，
    因此这里的日期列是月末有效性的主判据。只扫描物料表头之前的几行，
    避免把明细数据中的日期误当成库存快照日期。
    """
    cutoff_date = month_cutoff(month).date()
    parsed_dates: set[date] = set()
    max_row = max(1, header_row - 1)
    for row in ws.iter_rows(min_row=1, max_row=max_row, values_only=True):
        for value in row:
            parsed = parse_datetime(value)
            if parsed is not None:
                parsed_dates.add(parsed.date())
    ordered_dates = sorted(parsed_dates)
    return {
        "inventory_header_dates": ",".join(item.isoformat() for item in ordered_dates),
        "inventory_header_last_date": ordered_dates[-1].isoformat() if ordered_dates else "",
        "inventory_month_end_date": cutoff_date.isoformat(),
        "inventory_header_date_match": cutoff_date in parsed_dates,
        "inventory_header_warning": "" if cutoff_date in parsed_dates else "库存表表头上方未发现目标月末日期",
    }


def read_month_workbook(source: MonthlySource) -> WorkbookReadResult:
    wb = openpyxl.load_workbook(source.path, read_only=True, data_only=True)
    try:
        if "库存表" not in wb.sheetnames:
            raise ValueError(f"缺少库存表: {source.path}")
        ws = wb["库存表"]
        header_row, headers = _find_header(ws, set(REQUIRED_WORKBOOK_COLUMNS))
        code_column = headers.get("物料编码", headers.get("编号"))
        if code_column is None:
            raise ValueError(f"库存表缺少物料编码/编号: {source.path}")

        columns = {
            "code": code_column,
            "short_code": headers.get("编号"),
            "name": headers.get("物料名称"),
            "unit": headers.get("单位"),
            "inbound": headers["外仓入库总量"],
            "outbound": headers["外仓出库总量"],
            "stock": headers["库存"],
            "qualified_stock": headers.get("合格仓库存"),
            "external_required": headers.get("外应存"),
            "monthly_plan": headers.get("月计划"),
        }
        rows: OrderedDict[str, dict] = OrderedDict()
        duplicate_rows = 0
        for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
            full_code = normalize_code(row[columns["code"]] if columns["code"] < len(row) else None)
            short_code = normalize_code(row[columns["short_code"]] if columns["short_code"] is not None and columns["short_code"] < len(row) else None)
            if not full_code and not short_code:
                continue
            key = full_code or f"编号:{short_code}"
            record = {
                "month": source.month,
                "code": full_code or short_code,
                "short_code": short_code,
                "name": normalize_text(row[columns["name"]]) if columns["name"] is not None and columns["name"] < len(row) else "",
                "unit": normalize_text(row[columns["unit"]]) if columns["unit"] is not None and columns["unit"] < len(row) else "",
                "inbound": to_number(row[columns["inbound"]]),
                "outbound": to_number(row[columns["outbound"]]),
                "month_end_stock": to_number(row[columns["stock"]]),
                "qualified_stock": to_number(row[columns["qualified_stock"]]) if columns["qualified_stock"] is not None and columns["qualified_stock"] < len(row) else 0.0,
                "external_required": to_number(row[columns["external_required"]]) if columns["external_required"] is not None and columns["external_required"] < len(row) else 0.0,
                "monthly_plan": to_number(row[columns["monthly_plan"]]) if columns["monthly_plan"] is not None and columns["monthly_plan"] < len(row) else 0.0,
                "source_file": source.path.name,
            }
            if key in rows:
                duplicate_rows += 1
                for field in ("inbound", "outbound", "month_end_stock", "qualified_stock", "external_required", "monthly_plan"):
                    rows[key][field] += record[field]
            else:
                rows[key] = record
        quality = {
            "month": source.month,
            "source": source.source,
            "file": str(source.path),
            "status": "ok",
            "material_count": len(rows),
            "duplicate_material_rows": duplicate_rows,
            "missing_optional_columns": ",".join(column for column in OPTIONAL_WORKBOOK_COLUMNS if column not in headers),
        }
        quality.update(_read_inventory_header_quality(ws, header_row, source.month))
        quality.update(_read_detail_quality(wb, source.month))
        movement_records = _read_movement_rows(wb, source.month)
        if not quality["inventory_header_date_match"]:
            quality["status"] = "error"
            quality["error"] = quality["inventory_header_warning"]
        quality["movement_row_count"] = len(movement_records)
        return WorkbookReadResult(source.month, source.path, list(rows.values()), quality, movement_records)
    finally:
        wb.close()


def aggregate_sources(sources: Sequence[MonthlySource]) -> AggregationResult:
    if not sources:
        raise ValueError("没有可汇总的月末文件")
    records: list[dict] = []
    quality: list[dict] = []
    movement_records: list[dict] = []
    for source in sources:
        try:
            result = read_month_workbook(source)
        except Exception as exc:
            quality.append({"month": source.month, "source": source.source, "file": str(source.path), "status": "error", "error": str(exc)})
            continue
        quality.append(result.quality)
        if result.quality.get("status") == "error":
            continue
        records.extend(result.records)
        movement_records.extend(result.movement_records)
    if not records:
        errors = "; ".join(str(item.get("error", item.get("file", ""))) for item in quality)
        raise ValueError(f"没有可汇总数据: {errors}")
    return AggregationResult(records=records, quality=quality, sources=list(sources), movement_records=movement_records)


def select_latest_candidate(candidates: Iterable[MailAttachmentCandidate], cutoff: datetime) -> MailAttachmentCandidate:
    eligible = [candidate for candidate in candidates if candidate.received_at <= cutoff]
    if not eligible:
        raise ValueError(f"截止 {cutoff.isoformat()} 前没有符合条件的目标邮件")
    return max(eligible, key=lambda candidate: candidate.received_at)


def _imap_date(value: date) -> str:
    return value.strftime("%d-%b-%Y")


def _parse_message_candidates(
    *, uid: bytes, raw: bytes, metadata: bytes, subject_keyword: str, attachment_pattern: str
) -> list[MailAttachmentCandidate]:
    message = message_from_bytes(raw)
    received_at = parse_imap_internaldate(metadata) or parse_mail_datetime(message.get("Date"))
    if received_at is None:
        return []
    subject = decode_mime_header(message.get("Subject"))
    if subject_keyword not in subject:
        return []
    html = ""
    for part in message.walk():
        if part.get_content_type() != "text/html" or part.get_filename():
            continue
        payload = part.get_payload(decode=True)
        if payload is None:
            continue
        charset = part.get_content_charset() or "utf-8"
        html = payload.decode(charset, errors="replace")
        break
    candidates: list[MailAttachmentCandidate] = []
    for part in message.walk():
        filename = decode_mime_header(part.get_filename())
        if not filename or not fnmatch.fnmatchcase(filename, attachment_pattern):
            continue
        payload = part.get_payload(decode=True)
        if payload is None:
            continue
        candidates.append(MailAttachmentCandidate(uid.decode(errors="replace"), received_at, subject, filename, payload, html))
    return candidates


def fetch_latest_target_attachment(
    *,
    server: str,
    user: str,
    password: str,
    mailbox: str,
    month: str,
    subject_keyword: str,
    attachment_pattern: str,
) -> MailAttachmentCandidate:
    """只读 IMAP，按邮件时间选择目标月份截止前最新附件。"""
    cutoff = month_cutoff(month)
    candidates: list[MailAttachmentCandidate] = []
    mail = imaplib.IMAP4_SSL(server)
    try:
        mail.login(user, password)
        status, _ = mail.select(mailbox, readonly=True)
        if status != "OK":
            raise RuntimeError(f"无法只读打开邮箱目录: {mailbox}")
        status, data = mail.search(None, "SINCE", _imap_date(month_start(month)), "BEFORE", _imap_date(next_month_start(month)))
        if status != "OK":
            raise RuntimeError("IMAP 搜索失败")
        for uid in data[0].split():
            status, fetched = mail.fetch(uid, "(INTERNALDATE BODY.PEEK[])")
            if status != "OK":
                continue
            metadata = b" ".join(item[0] for item in fetched if isinstance(item, tuple) and isinstance(item[0], bytes))
            raw = next((item[1] for item in fetched if isinstance(item, tuple) and isinstance(item[1], bytes)), None)
            if raw is None:
                continue
            candidates.extend(_parse_message_candidates(
                uid=uid,
                raw=raw,
                metadata=metadata,
                subject_keyword=subject_keyword,
                attachment_pattern=attachment_pattern,
            ))
    finally:
        try:
            mail.logout()
        except Exception:
            pass
    return select_latest_candidate(candidates, cutoff)


def fetch_latest_target_attachment_any(
    *,
    server: str,
    user: str,
    password: str,
    mailbox: str,
    subject_keyword: str,
    attachment_pattern: str,
    limit: int = 100,
) -> MailAttachmentCandidate:
    """只读 IMAP，在最近若干封邮件中选择最新目标附件。"""
    candidates: list[MailAttachmentCandidate] = []
    mail = imaplib.IMAP4_SSL(server)
    try:
        mail.login(user, password)
        status, _ = mail.select(mailbox, readonly=True)
        if status != "OK":
            raise RuntimeError(f"无法只读打开邮箱目录: {mailbox}")
        status, data = mail.search(None, "ALL")
        if status != "OK":
            raise RuntimeError("IMAP 搜索失败")
        for uid in data[0].split()[-max(1, limit):]:
            status, fetched = mail.fetch(uid, "(INTERNALDATE BODY.PEEK[])")
            if status != "OK":
                continue
            metadata = b" ".join(item[0] for item in fetched if isinstance(item, tuple) and isinstance(item[0], bytes))
            raw = next((item[1] for item in fetched if isinstance(item, tuple) and isinstance(item[1], bytes)), None)
            if raw is not None:
                candidates.extend(_parse_message_candidates(
                    uid=uid,
                    raw=raw,
                    metadata=metadata,
                    subject_keyword=subject_keyword,
                    attachment_pattern=attachment_pattern,
                ))
    finally:
        try:
            mail.logout()
        except Exception:
            pass
    if not candidates:
        raise ValueError("最近邮件中没有符合条件的目标邮件")
    return max(candidates, key=lambda candidate: candidate.received_at)


def sha256_bytes(payload: bytes) -> str:
    return hashlib.sha256(payload).hexdigest()


def safe_filename(value: str) -> str:
    original = Path(value)
    stem = re.sub(r"[^A-Za-z0-9._-]+", "_", original.stem).strip("._") or "attachment"
    suffix = original.suffix if original.suffix.lower() in {".xlsx", ".xls"} else ".xlsx"
    return f"{stem}{suffix}"


def save_month_snapshot(
    *,
    candidate: MailAttachmentCandidate,
    month: str,
    auto_archive_dir: Path,
    raw_snapshot_dir: Path,
    manifest_path: Path | None = None,
) -> Path:
    auto_archive_dir.mkdir(parents=True, exist_ok=True)
    raw_snapshot_dir.mkdir(parents=True, exist_ok=True)
    digest = sha256_bytes(candidate.payload)
    stamp = candidate.received_at.astimezone(TZ_SHANGHAI).strftime("%Y%m%dT%H%M%S")
    raw_path = raw_snapshot_dir / f"{month_token(month)}_{stamp}_{safe_filename(candidate.filename)}"
    if not raw_path.exists():
        raw_path.write_bytes(candidate.payload)

    canonical_path = auto_archive_dir / f"{month_token(month)}月底.xlsx"
    if not canonical_path.exists() or sha256_bytes(canonical_path.read_bytes()) != digest:
        with tempfile.NamedTemporaryFile(dir=auto_archive_dir, prefix=".month-end-", suffix=".tmp", delete=False) as temp:
            temp.write(candidate.payload)
            temp_path = Path(temp.name)
        os.replace(temp_path, canonical_path)

    if manifest_path is None:
        return canonical_path
    append_manifest_record(
        manifest_path=manifest_path,
        candidate=candidate,
        month=month,
        raw_path=raw_path,
        canonical_path=canonical_path,
    )
    return canonical_path


def save_latest_mail(
    *,
    candidate: MailAttachmentCandidate,
    output_dir: Path,
) -> tuple[Path, Path | None]:
    """保存最新邮件附件和 HTML，供人工查看；不覆盖已有文件。"""
    output_dir.mkdir(parents=True, exist_ok=True)
    stamp = candidate.received_at.astimezone(TZ_SHANGHAI).strftime("%Y%m%d_%H%M%S")
    safe_name = safe_filename(candidate.filename)
    safe_path = Path(safe_name)
    attachment_path = output_dir / f"{safe_path.stem}_{stamp}{safe_path.suffix}"
    if not attachment_path.exists():
        attachment_path.write_bytes(candidate.payload)
    html_path: Path | None = None
    if candidate.html:
        html_path = output_dir / f"物料情况_{stamp}.html"
        if not html_path.exists():
            html_path.write_text(candidate.html, encoding="utf-8")
    return attachment_path, html_path


def append_manifest_record(
    *,
    manifest_path: Path,
    candidate: MailAttachmentCandidate,
    month: str,
    raw_path: Path,
    canonical_path: Path,
    quality: dict | None = None,
) -> None:
    manifest_path.parent.mkdir(parents=True, exist_ok=True)
    digest = sha256_bytes(candidate.payload)
    manifest_record = {
        "month": month,
        "uid": candidate.uid,
        "received_at": candidate.received_at.isoformat(),
        "subject": candidate.subject,
        "filename": candidate.filename,
        "sha256": digest,
        "raw_snapshot": str(raw_path),
        "canonical_file": str(canonical_path),
    }
    if quality:
        manifest_record.update({
            "detail_last_date": quality.get("detail_last_date", ""),
            "month_end_rows": quality.get("month_end_rows", 0),
            "future_rows": quality.get("future_rows", 0),
            "inventory_header_dates": quality.get("inventory_header_dates", ""),
            "inventory_header_last_date": quality.get("inventory_header_last_date", ""),
            "inventory_month_end_date": quality.get("inventory_month_end_date", ""),
            "inventory_header_date_match": quality.get("inventory_header_date_match", False),
            "movement_row_count": quality.get("movement_row_count", 0),
            "quality_status": quality.get("status", ""),
            "quality_warning": quality.get("detail_warning", ""),
        })
    with manifest_path.open("a", encoding="utf-8") as handle:
        handle.write(json.dumps(manifest_record, ensure_ascii=False) + "\n")


def _with_metadata(frame: pd.DataFrame, records: list[dict]) -> pd.DataFrame:
    if frame.empty:
        return frame
    metadata = pd.DataFrame(records)[["code", "short_code", "name", "unit"]].drop_duplicates("code")
    frame = frame.reset_index().rename(columns={"index": "code"})
    merged = metadata.merge(frame, on="code", how="right")
    return merged.rename(columns={"code": "物料编码", "short_code": "编号", "name": "物料名称", "unit": "单位"})


def build_report_frames(result: AggregationResult) -> OrderedDict[str, pd.DataFrame]:
    detail = pd.DataFrame(result.records)
    months = sorted(detail["month"].unique())
    detail_columns = [
        "month", "code", "short_code", "name", "unit", "inbound", "outbound",
        "month_end_stock", "qualified_stock", "external_required", "monthly_plan", "source_file",
    ]
    detail_frame = detail[detail_columns].sort_values(["month", "code"]).rename(columns={
        "month": "月份", "code": "物料编码", "short_code": "编号", "name": "物料名称", "unit": "单位",
        "inbound": "外仓入库总量", "outbound": "外仓出库总量", "month_end_stock": "月末库存",
        "qualified_stock": "合格仓库存", "external_required": "外应存", "monthly_plan": "月计划", "source_file": "来源文件",
    })

    summary_rows: list[dict] = []
    for code, group in detail.groupby("code", sort=True):
        latest = group.sort_values("month").iloc[-1]
        summary_rows.append({
            "code": code,
            "short_code": latest["short_code"],
            "name": latest["name"],
            "unit": latest["unit"],
            "累计入库": group["inbound"].sum(),
            "累计出库": group["outbound"].sum(),
            "累计净变化": group["inbound"].sum() - group["outbound"].sum(),
            "月均入库": group["inbound"].mean(),
            "月均出库": group["outbound"].mean(),
            "有业务月份数": int(((group["inbound"] != 0) | (group["outbound"] != 0)).sum()),
            "最新月份": latest["month"],
            "最新库存": latest["month_end_stock"],
            "最新合格仓库存": latest["qualified_stock"],
            "最新外应存": latest["external_required"],
            "最新月计划": latest["monthly_plan"],
        })
    overview = pd.DataFrame(summary_rows).sort_values(["累计出库", "code"], ascending=[False, True]).rename(columns={
        "code": "物料编码", "short_code": "编号", "name": "物料名称", "unit": "单位",
    })

    def wide(value: str, name: str) -> pd.DataFrame:
        pivot = detail.pivot_table(index="code", columns="month", values=value, aggfunc="sum", fill_value=0)
        pivot = pivot.reindex(columns=months, fill_value=0)
        pivot.columns = [f"{month_token(month)}{name}" for month in pivot.columns]
        return _with_metadata(pivot, result.records)

    quality = pd.DataFrame(result.quality)
    movement = pd.DataFrame(result.movement_records)
    if movement.empty:
        movement = pd.DataFrame(columns=[
            "month", "year", "date", "code", "name", "unit", "category",
            "business_type", "company", "note", "inbound", "outbound", "quantity",
        ])
    movement_detail = movement.rename(columns={
        "month": "月份", "year": "年度", "date": "日期", "code": "物料编码", "name": "物料名称",
        "unit": "单位", "category": "原始变动类别", "business_type": "业务类型", "company": "公司",
        "note": "备注", "inbound": "入库数量", "outbound": "出库数量", "quantity": "业务数量",
    })
    movement_detail = movement_detail[[
        "月份", "年度", "日期", "公司", "业务类型", "原始变动类别", "物料编码", "物料名称",
        "单位", "入库数量", "出库数量", "业务数量", "备注",
    ]].sort_values(["月份", "公司", "业务类型", "物料编码"])

    usage = movement[movement["business_type"] == "领用出库"].copy()
    usage_group_columns = ["month", "company", "code", "name", "unit"]
    if usage.empty:
        monthly_usage = pd.DataFrame(columns=["月份", "公司", "物料编码", "物料名称", "单位", "领用出库数量", "出库笔数"])
        annual_usage = pd.DataFrame(columns=["年度", "公司", "物料编码", "物料名称", "单位", "年度领用出库数量", "出库笔数"])
        company_month = pd.DataFrame()
        company_year = pd.DataFrame()
    else:
        monthly_usage = usage.groupby(usage_group_columns, as_index=False).agg(
            领用出库数量=("outbound", "sum"),
            出库笔数=("outbound", "size"),
        ).rename(columns={"month": "月份", "company": "公司", "code": "物料编码", "name": "物料名称", "unit": "单位"})
        monthly_usage = monthly_usage.sort_values(["月份", "领用出库数量", "公司", "物料编码"], ascending=[True, False, True, True])
        annual_usage = usage.groupby(["year", "company", "code", "name", "unit"], as_index=False).agg(
            年度领用出库数量=("outbound", "sum"),
            出库笔数=("outbound", "size"),
        ).rename(columns={"year": "年度", "company": "公司", "code": "物料编码", "name": "物料名称", "unit": "单位"})
        annual_usage = annual_usage.sort_values(["年度", "年度领用出库数量", "公司", "物料编码"], ascending=[True, False, True, True])
        company_month = usage.groupby(["month", "company"], as_index=False)["outbound"].sum()
        company_year = usage.groupby(["year", "company"], as_index=False)["outbound"].sum()

    business_summary = movement.groupby(["month", "company", "business_type"], as_index=False).agg(
        入库数量=("inbound", "sum"),
        出库数量=("outbound", "sum"),
        业务数量=("quantity", "sum"),
        笔数=("quantity", "size"),
    ).rename(columns={"month": "月份", "company": "公司", "business_type": "业务类型"})
    business_summary["年度"] = business_summary["月份"].str[:4]
    business_summary = business_summary[["月份", "年度", "公司", "业务类型", "入库数量", "出库数量", "业务数量", "笔数"]].sort_values(["月份", "公司", "业务类型"])

    def trend_frame(grouped: pd.DataFrame, index_column: str, value_name: str) -> pd.DataFrame:
        if grouped.empty:
            return pd.DataFrame(columns=[index_column])
        pivot = grouped.pivot_table(index=index_column, columns="company", values="outbound", aggfunc="sum", fill_value=0)
        totals = pivot.sum(axis=0).sort_values(ascending=False)
        pivot = pivot.loc[:, list(totals.head(8).index)]
        pivot = pivot.reset_index()
        pivot.columns.name = None
        pivot = pivot.rename(columns={"company": "公司"})
        return pivot

    company_month_trend = trend_frame(company_month, "month", "领用出库数量").rename(columns={"month": "月份"})
    company_year_trend = trend_frame(company_year, "year", "年度领用出库数量").rename(columns={"year": "年度"})
    type_trend = movement.groupby(["month", "business_type"], as_index=False)["quantity"].sum()
    if type_trend.empty:
        type_trend_frame = pd.DataFrame(columns=["月份"])
    else:
        type_trend_frame = type_trend.pivot_table(index="month", columns="business_type", values="quantity", aggfunc="sum", fill_value=0).reset_index()
        type_trend_frame.columns.name = None
        type_trend_frame = type_trend_frame.rename(columns={"month": "月份"})

    frames: OrderedDict[str, pd.DataFrame] = OrderedDict()
    frames["物料总览"] = overview
    frames["入库汇总"] = wide("inbound", "外仓入库总量")
    frames["出库汇总"] = wide("outbound", "外仓出库总量")
    frames["库存趋势"] = wide("month_end_stock", "月末库存")
    plan_detail = detail[["month", "code", "short_code", "name", "unit", "monthly_plan", "external_required", "month_end_stock", "source_file"]].copy()
    plan_detail = plan_detail.rename(columns={
        "month": "月份", "code": "物料编码", "short_code": "编号", "name": "物料名称", "unit": "单位",
        "monthly_plan": "月计划", "external_required": "外应存", "month_end_stock": "月末库存", "source_file": "来源文件",
    })
    frames["计划与库存"] = plan_detail.sort_values(["月份", "物料编码"])
    frames["物料明细"] = detail_frame
    frames["数据质量"] = quality
    frames["公司月度领用"] = monthly_usage
    frames["公司年度领用"] = annual_usage
    frames["公司业务汇总"] = business_summary
    frames["公司业务明细"] = movement_detail
    frames["公司领用趋势"] = company_month_trend
    frames["公司年度趋势"] = company_year_trend
    frames["业务类型趋势"] = type_trend_frame
    frames["图表"] = pd.DataFrame({"说明": [
        "领用出库按备注中的公司统一归类；借用、归还、退货、调拨等业务请查看公司业务汇总和公司业务明细。",
        "图表默认展示领用出库量最高的前 8 家公司。",
    ]})
    return frames


def write_report(frames: OrderedDict[str, pd.DataFrame], output_path: Path) -> Path:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        for sheet_name, frame in frames.items():
            frame.to_excel(writer, index=False, sheet_name=sheet_name)
    wb = openpyxl.load_workbook(output_path)
    try:
        for ws in wb.worksheets:
            ws.freeze_panes = "A2"
            if ws.max_row >= 1 and ws.max_column >= 1:
                ws.auto_filter.ref = f"A1:{openpyxl.utils.get_column_letter(ws.max_column)}{ws.max_row}"
            for column_cells in ws.columns:
                letter = column_cells[0].column_letter
                width = min(max(len(str(cell.value or "")) for cell in column_cells) + 2, 42)
                ws.column_dimensions[letter].width = max(width, 10)
        chart_ws = wb["图表"] if "图表" in wb.sheetnames else wb.create_sheet("图表")

        def add_line_chart(source_name: str, title: str, anchor: str) -> None:
            if source_name not in wb.sheetnames:
                return
            source_ws = wb[source_name]
            if source_ws.max_row < 2 or source_ws.max_column < 2:
                return
            chart = LineChart()
            chart.title = title
            chart.y_axis.title = "数量"
            chart.x_axis.title = "时间"
            chart.height = 8
            chart.width = 18
            data = Reference(source_ws, min_col=2, max_col=source_ws.max_column, min_row=1, max_row=source_ws.max_row)
            categories = Reference(source_ws, min_col=1, min_row=2, max_row=source_ws.max_row)
            chart.add_data(data, titles_from_data=True)
            chart.set_categories(categories)
            chart.style = 13
            chart.legend.position = "r"
            chart_ws.add_chart(chart, anchor)

        def add_bar_chart(source_name: str, title: str, anchor: str) -> None:
            if source_name not in wb.sheetnames:
                return
            source_ws = wb[source_name]
            if source_ws.max_row < 2 or source_ws.max_column < 2:
                return
            chart = BarChart()
            chart.type = "bar"
            chart.style = 10
            chart.title = title
            chart.y_axis.title = "公司"
            chart.x_axis.title = "数量"
            chart.height = 8
            chart.width = 18
            data = Reference(source_ws, min_col=2, max_col=source_ws.max_column, min_row=1, max_row=source_ws.max_row)
            categories = Reference(source_ws, min_col=1, min_row=2, max_row=source_ws.max_row)
            chart.add_data(data, titles_from_data=True)
            chart.set_categories(categories)
            chart.legend.position = "r"
            chart_ws.add_chart(chart, anchor)

        add_line_chart("公司领用趋势", "各公司每月领用出库趋势（前八）", "A4")
        add_bar_chart("公司年度趋势", "各公司年度领用出库趋势（前八）", "J4")
        add_line_chart("业务类型趋势", "各类业务数量趋势", "A22")
        wb.save(output_path)
    finally:
        wb.close()
    return output_path

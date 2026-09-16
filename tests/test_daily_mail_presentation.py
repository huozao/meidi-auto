import importlib.util
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook


ROOT = Path(__file__).resolve().parents[1]


def load_script(name: str, filename: str):
    path = ROOT / "script" / filename
    spec = importlib.util.spec_from_file_location(name, path)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"无法加载脚本: {path}")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


class DailyMailPresentationTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.email_script = load_script("send_email", "051 Send an email.py")
        cls.image_script = load_script("image_preview", "050 image.py")

    def test_subject_uses_inventory_business_date(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "总库存20260915.xlsx"
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.title = "库存表"
            worksheet["H3"] = "2026年09月15日"
            workbook.save(path)
            workbook.close()

            self.assertEqual(
                self.email_script.build_subject(str(path)),
                "美的库存及出入库日报｜2026-09-15｜库存、出入库及月计划",
            )

    def test_render_copy_only_keeps_inventory_sheet_visible(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            source = Path(temp_dir) / "source.xlsx"
            rendered = Path(temp_dir) / "render.xlsx"
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.title = "库存表"
            worksheet["A1"] = "标题"
            worksheet["A4"] = "表头"
            worksheet["A5"] = 1
            workbook.create_sheet("其他")
            workbook.save(source)
            workbook.close()

            bounds = self.image_script._prepare_render_workbook(
                source,
                rendered,
                col_start=1,
                col_end=20,
            )
            self.assertEqual(bounds, (1, 5))
            result = load_workbook(rendered, read_only=False, data_only=False)
            try:
                self.assertEqual(result["库存表"].print_area, "'库存表'!$A$1:$T$5")
                self.assertEqual(result["其他"].sheet_state, "hidden")
            finally:
                result.close()


if __name__ == "__main__":
    unittest.main()

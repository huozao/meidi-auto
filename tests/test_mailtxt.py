import importlib.util
import json
import tempfile
import unittest
from datetime import datetime
from pathlib import Path

import openpyxl


MODULE_PATH = Path(__file__).parents[1] / "script" / "050 mailtxt.py"
SPEC = importlib.util.spec_from_file_location("mailtxt", MODULE_PATH)
MAILTXT = importlib.util.module_from_spec(SPEC)
assert SPEC.loader is not None
SPEC.loader.exec_module(MAILTXT)


class MailDateHtmlTests(unittest.TestCase):
    def test_html_keeps_seconds_for_inventory_times(self):
        wb = openpyxl.Workbook()
        sheet = wb.active
        sheet["H3"] = datetime(2026, 9, 18, 8, 5, 6)
        sheet["M3"] = "2026-09-18T09:07:08+08:00"

        date, date2 = MAILTXT.get_dates(sheet)
        html = MAILTXT.construct_html_content(
            sheet, [], date, date2, 0, 0, 0, 0, 0, 0, 0
        )

        self.assertIn("2026-09-18 08:05:06", html)
        self.assertIn("2026-09-18 09:07:08", html)
        self.assertNotIn("2026-09-18</strong>", html)

    def test_html_prefers_source_mail_received_times(self):
        wb = openpyxl.Workbook()
        sheet = wb.active
        sheet["H3"] = "2026年09月18日"
        sheet["M3"] = "2026-09-18"
        with tempfile.TemporaryDirectory() as temp_dir:
            meta_path = Path(temp_dir) / "mail_meta.json"
            meta_path.write_text(json.dumps({
                "selected_heyu_da_received_at": "2026-09-18T10:11:12+08:00",
                "selected_waiting_received_at": "2026-09-18T13:14:15+08:00",
            }), encoding="utf-8")

            self.assertEqual(
                MAILTXT.get_dates(sheet, meta_path),
                ("2026-09-18 10:11:12", "2026-09-18 13:14:15"),
            )

    def test_outdated_home_stock_is_styled_red(self):
        wb = openpyxl.Workbook()
        sheet = wb.active
        # 俊都仓储为 10-10，家里库存停留在 09-30
        date = "2026-10-10 21:06:25"
        date2 = "2026-09-30 18:00:05"

        html = MAILTXT.construct_html_content(
            sheet, [], date, date2, 0, 0, 0, 0, 0, 0, 0
        )

        # 家里库存卡片应该有醒目的红色字体、outdated 警示样式及非当天新鲜日期标注
        self.assertIn("2026-09-30 18:00:05", html)
        self.assertIn('class="date-card outdated"', html)
        self.assertIn('<span style="color: #e53e3e; font-weight: bold;">家里库存</span>', html)
        self.assertIn('<strong style="color: #e53e3e; font-weight: bold;">2026-09-30 18:00:05（非当天新鲜日期）</strong>', html)

    def test_is_same_day_helper(self):
        # 相同日期
        self.assertTrue(MAILTXT._is_same_day("2026-10-10 18:00:00", base_date_str="2026-10-10 21:00:00"))
        # 跨天不同日期
        self.assertFalse(MAILTXT._is_same_day("2026-09-30 18:00:05", base_date_str="2026-10-10 21:00:00"))
        # 异常缺失状态
        self.assertFalse(MAILTXT._is_same_day("未同步（缺失）", base_date_str="2026-10-10 21:00:00"))
        self.assertFalse(MAILTXT._is_same_day(None))


if __name__ == "__main__":
    unittest.main()

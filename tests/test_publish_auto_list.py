import io
import importlib.util
import unittest
from pathlib import Path

import openpyxl


MODULE_PATH = Path(__file__).parents[1] / "tools" / "publish_auto_list.py"
SPEC = importlib.util.spec_from_file_location("publish_auto_list", MODULE_PATH)
MODULE = importlib.util.module_from_spec(SPEC)
assert SPEC.loader is not None
SPEC.loader.exec_module(MODULE)


class PublishAutoListTests(unittest.TestCase):
    def _workbook(self, external, home, old=True, one_week=False):
        wb = openpyxl.Workbook()
        ws = wb.active
        old_external = "外应存(1周)" if one_week else "外应存(3月周均)"
        old_home = "家应存(1周)" if one_week else "家应存(3月周均)"
        external_header = old_external if old else "外应存(3月2周均)"
        home_header = old_home if old else "家应存(3月2周均)"
        ws.append(["编号", external_header, home_header])
        ws.append(["A", external, home])
        output = io.BytesIO()
        wb.save(output)
        wb.close()
        return output.getvalue()

    def test_migrates_old_values_and_headers(self):
        result, changed = MODULE.migrate_legacy_list(self._workbook(10, 12.5))
        wb = openpyxl.load_workbook(io.BytesIO(result), data_only=True)
        row = list(wb.active.iter_rows(values_only=True))
        wb.close()

        self.assertEqual(changed, 1)
        self.assertEqual(row[0], ("编号", "外应存(3月2周均)", "家应存(3月2周均)"))
        self.assertEqual(row[1], ("A", 20, 25))

    def test_refuses_to_migrate_new_list_twice(self):
        source = self._workbook(10, 12.5, old=False)
        with self.assertRaises(ValueError):
            MODULE.migrate_legacy_list(source)

    def test_migrates_one_week_headers(self):
        result, changed = MODULE.migrate_legacy_list(self._workbook(10, 12.5, one_week=True))
        wb = openpyxl.load_workbook(io.BytesIO(result), data_only=True)
        row = list(wb.active.iter_rows(values_only=True))
        wb.close()

        self.assertEqual(changed, 1)
        self.assertEqual(row[0][1:], ("外应存(3月2周均)", "家应存(3月2周均)"))
        self.assertEqual(row[1][1:], (20, 25))


if __name__ == "__main__":
    unittest.main()
